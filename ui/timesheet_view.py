import os
import json
from datetime import datetime, timedelta
import tkinter as tk
from tkinter import messagebox, filedialog
import customtkinter as ctk

from tracker.logger import logger
from analyzer.gemma_client import GemmaClient
from analyzer.json_parser import load_timesheet_from_file, save_timesheet_to_file

CATEGORY_COLORS = {
    "Development": "#3b82f6",
    "Documentation & Excel": "#10b981",
    "Communication & Email": "#8b5cf6",
    "Research & Web": "#06b6d4",
    "Admin & Forms": "#f59e0b",
    "General": "#64748b"
}

class TimesheetView(ctk.CTkFrame):
    def __init__(self, master, gemma_client=None, **kwargs):
        super().__init__(master, **kwargs)
        self.gemma_client = gemma_client or GemmaClient()
        self.current_date = datetime.now()

        self._setup_ui()
        self.load_timesheet_for_date(self.current_date.strftime("%Y-%m-%d"))

    def _setup_ui(self):
        # 1. Date Navigation Header
        self.header_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.header_frame.pack(fill="x", padx=15, pady=(15, 10))

        self.btn_prev_day = ctk.CTkButton(
            self.header_frame, text="< Prev Day", width=90, height=32,
            command=self._on_prev_day
        )
        self.btn_prev_day.pack(side="left")

        self.lbl_date = ctk.CTkLabel(
            self.header_frame, text="", font=ctk.CTkFont(size=18, weight="bold")
        )
        self.lbl_date.pack(side="left", expand=True)

        self.btn_today = ctk.CTkButton(
            self.header_frame, text="Today", width=70, height=32, fg_color="#334155",
            command=self._on_today
        )
        self.btn_today.pack(side="right", padx=(5, 0))

        self.btn_next_day = ctk.CTkButton(
            self.header_frame, text="Next Day >", width=90, height=32,
            command=self._on_next_day
        )
        self.btn_next_day.pack(side="right")

        # 2. Daily Summary Card
        self.summary_card = ctk.CTkFrame(self, corner_radius=10, fg_color="#1e293b")
        self.summary_card.pack(fill="x", padx=15, pady=8)

        self.lbl_summary_title = ctk.CTkLabel(
            self.summary_card, text="Daily Executive Summary", font=ctk.CTkFont(size=14, weight="bold"),
            text_color="#94a3b8"
        )
        self.lbl_summary_title.pack(anchor="w", padx=16, pady=(10, 4))

        self.lbl_summary_text = ctk.CTkLabel(
            self.summary_card, text="No timesheet generated yet for this date.",
            wraplength=760, justify="left", font=ctk.CTkFont(size=13)
        )
        self.lbl_summary_text.pack(anchor="w", padx=16, pady=(0, 10))

        # Stats bar inside summary card
        self.stats_bar = ctk.CTkFrame(self.summary_card, fg_color="transparent")
        self.stats_bar.pack(fill="x", padx=16, pady=(0, 12))

        self.lbl_total_hours = ctk.CTkLabel(
            self.stats_bar, text="Total Hours: --", font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#38bdf8"
        )
        self.lbl_total_hours.pack(side="left", padx=(0, 20))

        self.lbl_productive_hours = ctk.CTkLabel(
            self.stats_bar, text="Productive: --", font=ctk.CTkFont(size=13, weight="bold"),
            text_color="#4ade80"
        )
        self.lbl_productive_hours.pack(side="left")

        # 3. Category Breakdown Pills Frame
        self.categories_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.categories_frame.pack(fill="x", padx=15, pady=4)

        # 4. Scrollable Timesheet Entries
        self.entries_scroll = ctk.CTkScrollableFrame(self, corner_radius=8)
        self.entries_scroll.pack(fill="both", expand=True, padx=15, pady=8)

        # 5. Bottom Action Bar
        self.action_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.action_frame.pack(fill="x", padx=15, pady=(8, 15))

        self.btn_analyze = ctk.CTkButton(
            self.action_frame, text="Generate Timesheet with Gemma",
            height=38, font=ctk.CTkFont(size=13, weight="bold"), fg_color="#6366f1", hover_color="#4f46e5",
            command=self.analyze_current_date
        )
        self.btn_analyze.pack(side="left", padx=(0, 10))

        self.btn_export_excel = ctk.CTkButton(
            self.action_frame, text="Export Excel (.xlsx)", width=130, height=38,
            fg_color="#059669", hover_color="#047857", command=self.export_excel
        )
        self.btn_export_excel.pack(side="left", padx=(0, 8))

        self.btn_export_json = ctk.CTkButton(
            self.action_frame, text="Export JSON", width=100, height=38,
            fg_color="#334155", hover_color="#475569", command=self.export_json
        )
        self.btn_export_json.pack(side="left", padx=(0, 8))

        self.btn_copy_text = ctk.CTkButton(
            self.action_frame, text="Copy Report", width=100, height=38,
            fg_color="#334155", hover_color="#475569", command=self.copy_clipboard
        )
        self.btn_copy_text.pack(side="left")

        self.lbl_status = ctk.CTkLabel(self.action_frame, text="", text_color="#94a3b8")
        self.lbl_status.pack(side="right")

    def _on_prev_day(self):
        self.current_date -= timedelta(days=1)
        self.load_timesheet_for_date(self.current_date.strftime("%Y-%m-%d"))

    def _on_next_day(self):
        self.current_date += timedelta(days=1)
        self.load_timesheet_for_date(self.current_date.strftime("%Y-%m-%d"))

    def _on_today(self):
        self.current_date = datetime.now()
        self.load_timesheet_for_date(self.current_date.strftime("%Y-%m-%d"))

    def load_timesheet_for_date(self, date_str):
        formatted_date = datetime.strptime(date_str, "%Y-%m-%d").strftime("%A, %B %d, %Y")
        self.lbl_date.configure(text=formatted_date)
        self.current_timesheet = load_timesheet_from_file(date_str)

        # Clear existing entry widgets
        for widget in self.entries_scroll.winfo_children():
            widget.destroy()

        for widget in self.categories_frame.winfo_children():
            widget.destroy()

        if not self.current_timesheet:
            self.lbl_summary_text.configure(
                text="No timesheet generated yet for this date. Click 'Generate Timesheet with Gemma' to analyze today's activities."
            )
            self.lbl_total_hours.configure(text="Total Hours: --")
            self.lbl_productive_hours.configure(text="Productive: --")

            placeholder = ctk.CTkLabel(
                self.entries_scroll,
                text="No entries recorded yet.\nClick below to run Gemma analysis on today's activity log.",
                font=ctk.CTkFont(size=14), text_color="#64748b"
            )
            placeholder.pack(pady=40)
            return

        # Populate summary
        self.lbl_summary_text.configure(text=self.current_timesheet.get("summary", ""))
        total_h = self.current_timesheet.get("total_hours_tracked", 0)
        prod_h = self.current_timesheet.get("productive_hours", total_h)
        self.lbl_total_hours.configure(text=f"Total Hours: {total_h}h")
        self.lbl_productive_hours.configure(text=f"Productive: {prod_h}h")

        # Category pills
        cat_summary = self.current_timesheet.get("categories_summary", {})
        for cat_name, hours in cat_summary.items():
            color = CATEGORY_COLORS.get(cat_name, "#64748b")
            pill = ctk.CTkFrame(self.categories_frame, fg_color=color, corner_radius=12)
            pill.pack(side="left", padx=4, pady=2)
            lbl = ctk.CTkLabel(
                pill, text=f"{cat_name}: {hours}h", font=ctk.CTkFont(size=11, weight="bold"),
                text_color="#ffffff"
            )
            lbl.pack(padx=8, pady=2)

        # Entries
        entries = self.current_timesheet.get("entries", [])
        for entry in entries:
            self._render_entry_card(entry)

    def _render_entry_card(self, entry):
        card = ctk.CTkFrame(self.entries_scroll, corner_radius=8, fg_color="#1e293b")
        card.pack(fill="x", pady=5, padx=2)

        top_row = ctk.CTkFrame(card, fg_color="transparent")
        top_row.pack(fill="x", padx=12, pady=(10, 4))

        # Time range
        start = entry.get("start_time", "")
        end = entry.get("end_time", "")
        dur = entry.get("duration_minutes", 0)
        time_text = f"{start} - {end} ({dur}m)"
        lbl_time = ctk.CTkLabel(
            top_row, text=time_text, font=ctk.CTkFont(size=12, weight="bold"),
            text_color="#94a3b8"
        )
        lbl_time.pack(side="left")

        # Category Badge
        cat = entry.get("category", "General")
        cat_color = CATEGORY_COLORS.get(cat, "#64748b")
        cat_badge = ctk.CTkLabel(
            top_row, text=f"  {cat}  ", font=ctk.CTkFont(size=11, weight="bold"),
            fg_color=cat_color, text_color="#ffffff", corner_radius=6
        )
        cat_badge.pack(side="right")

        # Project Name
        project = entry.get("project", "General")
        lbl_proj = ctk.CTkLabel(
            card, text=project, font=ctk.CTkFont(size=14, weight="bold"),
            text_color="#f8fafc"
        )
        lbl_proj.pack(anchor="w", padx=12, pady=(2, 2))

        # Activity Description
        activity = entry.get("activity", "")
        lbl_act = ctk.CTkLabel(
            card, text=activity, wraplength=740, justify="left",
            font=ctk.CTkFont(size=12), text_color="#cbd5e1"
        )
        lbl_act.pack(anchor="w", padx=12, pady=(2, 6))

        # Artifacts (Files, URLs, Emails)
        artifacts = entry.get("artifacts", [])
        if artifacts:
            art_row = ctk.CTkFrame(card, fg_color="transparent")
            art_row.pack(fill="x", padx=12, pady=(0, 10))

            lbl_tag = ctk.CTkLabel(art_row, text="Artifacts: ", font=ctk.CTkFont(size=11, weight="bold"), text_color="#64748b")
            lbl_tag.pack(side="left")

            for art in artifacts:
                art_pill = ctk.CTkLabel(
                    art_row, text=f" {art} ", font=ctk.CTkFont(size=10),
                    fg_color="#0f172a", text_color="#38bdf8", corner_radius=4
                )
                art_pill.pack(side="left", padx=3)

    def analyze_current_date(self):
        date_str = self.current_date.strftime("%Y-%m-%d")
        raw_log = logger.read_log_text(date_str)

        if not raw_log.strip():
            messagebox.showinfo("No Activity Log", f"No activity log found for {date_str}. Make sure tracking is active.")
            return

        self.btn_analyze.configure(state="disabled", text="Gemma is analyzing logs...")
        self.lbl_status.configure(text="Connecting to LM Studio...")

        import threading
        def worker():
            client = self.gemma_client
            timesheet, msg = client.analyze_daily_log(date_str, raw_log)

            def on_done():
                self.btn_analyze.configure(state="normal", text="Generate Timesheet with Gemma")
                self.lbl_status.configure(text=msg)
                if timesheet:
                    self.load_timesheet_for_date(date_str)
                    messagebox.showinfo("Timesheet Ready", f"Analysis complete for {date_str}!\n\n{timesheet.get('summary')}")
                else:
                    messagebox.showerror("Analysis Error", msg)

            self.after(0, on_done)

        threading.Thread(target=worker, daemon=True).start()

    def export_excel(self):
        if not getattr(self, "current_timesheet", None):
            messagebox.showinfo("No Timesheet", "Please generate a timesheet first.")
            return

        date_str = self.current_date.strftime("%Y-%m-%d")
        default_name = f"timesheet_{date_str}.xlsx"
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile=default_name
        )
        if not file_path:
            return

        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"Timesheet {date_str}"

            # Styles
            title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
            data_font = Font(name="Calibri", size=10)
            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            # Title
            ws["A1"] = f"TIMESHEET REPORT - {date_str}"
            ws["A1"].font = title_font
            ws["A2"] = f"Summary: {self.current_timesheet.get('summary', '')}"
            ws["A2"].font = Font(italic=True, size=10, color="475569")
            ws["A3"] = f"Total Hours Tracked: {self.current_timesheet.get('total_hours_tracked', '')}h  |  Productive Hours: {self.current_timesheet.get('productive_hours', '')}h"
            ws["A3"].font = Font(bold=True, size=10, color="059669")

            # Table Headers
            headers = ["ID", "Start Time", "End Time", "Duration (min)", "Project", "Category", "Activity Description", "Artifacts"]
            row_idx = 5
            for col_idx, h in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Entries
            for item in self.current_timesheet.get("entries", []):
                row_idx += 1
                artifacts_str = ", ".join(item.get("artifacts", []))
                row_values = [
                    item.get("id", ""),
                    item.get("start_time", ""),
                    item.get("end_time", ""),
                    item.get("duration_minutes", ""),
                    item.get("project", ""),
                    item.get("category", ""),
                    item.get("activity", ""),
                    artifacts_str
                ]
                for col_idx, val in enumerate(row_values, 1):
                    c = ws.cell(row=row_idx, column=col_idx, value=val)
                    c.font = data_font
                    c.border = thin_border
                    if col_idx in [1, 2, 3, 4]:
                        c.alignment = Alignment(horizontal="center")

            # Adjust column widths
            col_widths = [6, 12, 12, 16, 22, 22, 45, 30]
            for i, width in enumerate(col_widths, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

            wb.save(file_path)
            messagebox.showinfo("Export Successful", f"Timesheet exported to Excel:\n{file_path}")

        except ImportError:
            # Fallback to CSV if openpyxl is missing
            csv_path = file_path.replace(".xlsx", ".csv")
            import csv
            with open(csv_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["ID", "Start Time", "End Time", "Duration (min)", "Project", "Category", "Activity Description", "Artifacts"])
                for item in self.current_timesheet.get("entries", []):
                    writer.writerow([
                        item.get("id", ""), item.get("start_time", ""), item.get("end_time", ""),
                        item.get("duration_minutes", ""), item.get("project", ""), item.get("category", ""),
                        item.get("activity", ""), ", ".join(item.get("artifacts", []))
                    ])
            messagebox.showinfo("Exported as CSV", f"openpyxl not installed. Saved as CSV instead:\n{csv_path}")

    def export_json(self):
        if not getattr(self, "current_timesheet", None):
            messagebox.showinfo("No Timesheet", "Please generate a timesheet first.")
            return

        date_str = self.current_date.strftime("%Y-%m-%d")
        file_path = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("JSON Files", "*.json")],
            initialfile=f"timesheet_{date_str}.json"
        )
        if file_path:
            with open(file_path, "w", encoding="utf-8") as f:
                json.dump(self.current_timesheet, f, indent=2)
            messagebox.showinfo("Saved", f"Timesheet JSON saved to:\n{file_path}")

    def copy_clipboard(self):
        if not getattr(self, "current_timesheet", None):
            messagebox.showinfo("No Timesheet", "Please generate a timesheet first.")
            return

        lines = [
            f"# Timesheet - {self.current_timesheet.get('date')}",
            f"**Total Hours:** {self.current_timesheet.get('total_hours_tracked')}h | **Productive:** {self.current_timesheet.get('productive_hours')}h",
            f"\n**Summary:** {self.current_timesheet.get('summary')}\n",
            "| Time | Project | Category | Activity |",
            "| --- | --- | --- | --- |"
        ]
        for e in self.current_timesheet.get("entries", []):
            time_str = f"{e.get('start_time')}-{e.get('end_time')}"
            lines.append(f"| {time_str} | {e.get('project')} | {e.get('category')} | {e.get('activity')} |")

        text = "\n".join(lines)
        self.clipboard_clear()
        self.clipboard_append(text)
        messagebox.showinfo("Copied", "Formatted timesheet copied to clipboard!")
